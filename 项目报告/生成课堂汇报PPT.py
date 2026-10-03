from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
ASSETS = Path(r"C:\Users\阿斯顿\.cursor\projects\d\assets")
OUT = Path(__file__).with_name("青禾门诊智能问诊系统-课堂汇报-BUG证据版.pptx")
PPT_ASSETS = Path(__file__).with_name("PPT素材")
PPT_ASSETS.mkdir(exist_ok=True)
BUG_ASSETS = Path(__file__).with_name("BUG证据")

W, H = Inches(13.333333), Inches(7.5)

TEAL = RGBColor(8, 124, 146)
MINT = RGBColor(18, 184, 166)
CYAN = RGBColor(35, 211, 195)
PALE = RGBColor(234, 251, 250)
PAPER = RGBColor(247, 252, 253)
NAVY = RGBColor(22, 59, 92)
TEXT = RGBColor(69, 103, 125)
MUTED = RGBColor(120, 148, 166)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(205, 69, 84)
ORANGE = RGBColor(231, 146, 72)
BLUE = RGBColor(63, 114, 190)
PURPLE = RGBColor(119, 99, 205)
GREEN = RGBColor(43, 153, 122)
LINE = RGBColor(214, 239, 239)

FONT = "Microsoft YaHei"
FONT_DISPLAY = "Microsoft YaHei"

ROBOT = ROOT / "commerce-cs-agent" / "web" / "public" / "brand" / "ai_robot.png"
ASK_SHOT_SOURCE = ASSETS / "qh-preview.png"
PROFILE_SHOT_SOURCE = ASSETS / (
    "c__Users_____AppData_Roaming_Cursor_User_workspaceStorage_"
    "eb80d56d9d0cf8300e47dabaf73dfb88_images___________-"
    "eb2ddccc-6e4a-45c6-9eda-6a6cc5833bb2.jpg"
)
KB_SHOT_SOURCE = ASSETS / (
    "c__Users_____AppData_Roaming_Cursor_User_workspaceStorage_"
    "eb80d56d9d0cf8300e47dabaf73dfb88_images____________-"
    "b713cded-3918-41e4-8815-4475a241a530.jpg"
)
RED_FLAG_SHOT = ASSETS / (
    "c__Users_____AppData_Roaming_Cursor_User_workspaceStorage_"
    "eb80d56d9d0cf8300e47dabaf73dfb88_images_image-6c300e95-"
    "9abc-4092-9f18-04d83476bcee.png"
)
BUG_1_SHOT = BUG_ASSETS / "Bug-01-阅读文章时搜索不生效.png"
BUG_2_SHOT = BUG_ASSETS / "Bug-02-感冒文章被分到用药指南.png"
BUG_3_SHOT = BUG_ASSETS / "Bug-03-新会话首条消息未应用就诊卡.png"


def crop_top(source, output_name, top):
    output = PPT_ASSETS / output_name
    with Image.open(source) as image:
        image.crop((0, top, image.width, image.height)).save(output)
    return output


# 旧截图顶部曾包含现已删除的“链路”导航；演示页仅保留实际功能区域。
ASK_SHOT = crop_top(ASK_SHOT_SOURCE, "问诊界面.png", 62)
PROFILE_SHOT = crop_top(PROFILE_SHOT_SOURCE, "就诊卡界面.jpg", 54)
KB_SHOT = crop_top(KB_SHOT_SOURCE, "知识库界面.jpg", 43)


def rgb_hex(color):
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def set_bg(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text, x, y, w, h, size=18, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, font=FONT, valign=MSO_ANCHOR.MIDDLE,
             margin=0.03, line_spacing=1.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_text(slide, runs, x, y, w, h, size=18, align=PP_ALIGN.LEFT,
                  valign=MSO_ANCHOR.MIDDLE, margin=0.03):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    p.alignment = align
    for text, color, bold, run_size in runs:
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(run_size or size)
        r.font.color.rgb = color
        r.font.bold = bold
    return box


def shape(slide, kind, x, y, w, h, fill=WHITE, line=LINE, radius=True, shadow=False):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(1)
    if shadow:
        # python-pptx does not expose full shadow formatting; a translucent backing
        # is added by callers where needed.
        pass
    return shp


def card(slide, x, y, w, h, fill=WHITE, line=LINE):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line)


def circle(slide, x, y, d, fill, line=None):
    return shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill, line or fill)


def line(slide, x1, y1, x2, y2, color=LINE, width=1.5, arrow=False):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    if arrow:
        conn.line.end_arrowhead = True
    return conn


def add_picture_crop(slide, path, x, y, w, h, rounded=False):
    path = Path(path)
    with Image.open(path) as img:
        iw, ih = img.size
    target = w / h
    source = iw / ih
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if source > target:
        crop = (1 - target / source) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    elif source < target:
        crop = (1 - source / target) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    return pic


def add_picture_contain(slide, path, x, y, w, h):
    path = Path(path)
    with Image.open(path) as img:
        iw, ih = img.size
    scale = min(w / iw, h / ih)
    draw_w = iw * scale
    draw_h = ih * scale
    return slide.shapes.add_picture(
        str(path),
        Inches(x + (w - draw_w) / 2),
        Inches(y + (h - draw_h) / 2),
        Inches(draw_w),
        Inches(draw_h),
    )


def add_brand(slide, inverse=False):
    heart = slide.shapes.add_shape(
        MSO_SHAPE.HEART, Inches(0.38), Inches(0.23), Inches(0.32), Inches(0.32)
    )
    heart.fill.solid()
    heart.fill.fore_color.rgb = WHITE if inverse else MINT
    heart.line.fill.background()
    add_text(
        slide, "青禾门诊", 0.75, 0.18, 1.8, 0.33, 13,
        WHITE if inverse else TEAL, True
    )


def add_footer(slide, index, label="青禾门诊 · 智能问诊 · 教学参考"):
    line(slide, 0.55, 7.12, 12.78, 7.12, LINE, 0.8)
    add_text(slide, label, 0.58, 7.16, 5.4, 0.18, 8, MUTED)
    add_text(slide, f"{index:02d}", 12.25, 7.14, 0.48, 0.2, 9, TEAL, True, PP_ALIGN.RIGHT)


def add_title(slide, index, eyebrow, title, subtitle=None):
    add_brand(slide)
    add_text(slide, f"{index:02d}  {eyebrow}", 0.63, 0.78, 3.2, 0.25, 10, MINT, True)
    add_text(slide, title, 0.62, 1.03, 11.7, 0.62, 27, NAVY, True, valign=MSO_ANCHOR.TOP)
    if subtitle:
        add_text(slide, subtitle, 0.64, 1.63, 11.6, 0.34, 11, MUTED, valign=MSO_ANCHOR.TOP)


def add_label(slide, text, x, y, w, fill=PALE, color=TEAL):
    shp = card(slide, x, y, w, 0.36, fill, fill)
    add_text(slide, text, x, y, w, 0.36, 10, color, True, PP_ALIGN.CENTER)
    return shp


def add_bullet_list(slide, items, x, y, w, h, size=14, color=TEXT,
                    bullet_color=MINT, gap=0.52):
    for i, item in enumerate(items):
        cy = y + i * gap
        circle(slide, x, cy + 0.12, 0.11, bullet_color)
        add_text(slide, item, x + 0.22, cy, w - 0.22, 0.38, size, color, valign=MSO_ANCHOR.TOP)


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, TEAL)
    # Soft abstract arcs
    for x, y, d, color in [
        (8.0, -2.0, 7.2, RGBColor(19, 154, 168)),
        (9.3, -0.9, 5.5, RGBColor(30, 177, 174)),
        (-1.2, 5.8, 3.2, RGBColor(25, 152, 159)),
    ]:
        circle(slide, x, y, d, color, color)
    add_brand(slide, inverse=True)
    add_label(slide, "软件工程课程 · 小组项目", 0.72, 1.22, 2.2, RGBColor(27, 156, 164), WHITE)
    add_text(slide, "青禾门诊", 0.73, 1.82, 6.6, 0.72, 38, WHITE, True, valign=MSO_ANCHOR.TOP)
    add_text(slide, "智能问诊系统", 0.73, 2.48, 6.6, 0.72, 38, WHITE, True, valign=MSO_ANCHOR.TOP)
    add_text(slide, "单 Agent · RAG 知识增强 · Java 安全门禁", 0.78, 3.42, 6.4, 0.42, 17, PALE)
    add_text(
        slide, "教学演示，不替代面诊，不生成真实处方", 0.78, 4.02, 5.9, 0.35,
        12, RGBColor(205, 244, 242)
    )
    card(slide, 0.75, 5.15, 5.5, 1.15, RGBColor(16, 135, 151), RGBColor(50, 177, 184))
    add_text(slide, "小组：________________", 1.02, 5.43, 2.4, 0.32, 12, WHITE)
    add_text(slide, "成员：________________", 3.45, 5.43, 2.4, 0.32, 12, WHITE)
    add_text(slide, "2026 · 课堂汇报", 1.02, 5.88, 4.5, 0.25, 10, RGBColor(198, 241, 239))
    # Robot on white medallion
    circle(slide, 8.03, 0.92, 4.62, RGBColor(244, 254, 253), RGBColor(176, 231, 228))
    slide.shapes.add_picture(str(ROBOT), Inches(8.36), Inches(1.18), Inches(4.0), Inches(4.0))
    add_text(slide, "AI 让医疗更有温度", 8.47, 5.47, 3.8, 0.38, 16, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, "QINGHE CLINIC", 8.77, 5.93, 3.2, 0.25, 9, RGBColor(190, 239, 236), True, PP_ALIGN.CENTER)
    return slide


def slide_background(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 2, "WHY", "为什么做青禾门诊？", "从通用聊天走向“有知识、有门禁、可回归”的问诊系统")
    # Left problem statement
    card(slide, 0.63, 2.1, 4.0, 4.45, WHITE, LINE)
    add_label(slide, "现实问题", 0.9, 2.38, 1.1, RGBColor(255, 242, 242), RED)
    add_text(slide, "通用聊天模型\n能回答，但不一定安全", 0.9, 2.92, 3.3, 0.94, 23, NAVY, True, valign=MSO_ANCHOR.TOP)
    add_bullet_list(
        slide,
        ["用户往往无法一次说清病情", "回答可能缺少权威依据", "危急症状不能继续“给方案”", "模型或检索宕机会打断演示"],
        0.95, 4.15, 3.3, 1.9, 13
    )
    # Right answer cards
    add_text(slide, "我们的回答", 5.05, 2.22, 3.2, 0.35, 17, TEAL, True)
    specs = [
        ("01", "多轮问诊", "主动追问时长、部位、发热等关键槽位", MINT),
        ("02", "知识增强", "RAG 取回资料后再组织回答，并显示出处", BLUE),
        ("03", "安全门禁", "红旗只导向急诊；药品必须经过白名单", RED),
        ("04", "完整系统", "就诊卡、历史、知识库、语音与图片协同", PURPLE),
    ]
    positions = [(5.02, 2.72), (8.93, 2.72), (5.02, 4.53), (8.93, 4.53)]
    for (num, title, desc, color), (x, y) in zip(specs, positions):
        card(slide, x, y, 3.45, 1.48, WHITE, LINE)
        circle(slide, x + 0.22, y + 0.22, 0.44, color)
        add_text(slide, num, x + 0.22, y + 0.22, 0.44, 0.44, 10, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, title, x + 0.8, y + 0.18, 2.25, 0.34, 16, NAVY, True)
        add_text(slide, desc, x + 0.8, y + 0.58, 2.3, 0.62, 11, TEXT, valign=MSO_ANCHOR.TOP)
    add_footer(slide, 2)
    return slide


def slide_requirements(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 3, "SCOPE", "用户、用例与范围边界", "先把“做什么”与“不做什么”说清楚")
    # Roles
    roles = [
        ("演示患者", "文字 / 语音 / 图片描述不适\n查看建议与历史", MINT),
        ("医学编辑", "浏览、检索和补充知识库\n维护教学资料", BLUE),
        ("答辩评委", "观察主链路、安全行为\n验证系统工作量", PURPLE),
    ]
    for i, (title, desc, color) in enumerate(roles):
        x = 0.65 + i * 4.18
        card(slide, x, 2.05, 3.78, 1.42, WHITE, LINE)
        circle(slide, x + 0.25, 2.31, 0.66, color)
        add_text(slide, str(i + 1), x + 0.25, 2.31, 0.66, 0.66, 15, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, title, x + 1.08, 2.18, 2.25, 0.34, 15, NAVY, True)
        add_text(slide, desc, x + 1.08, 2.58, 2.3, 0.58, 10.5, TEXT, valign=MSO_ANCHOR.TOP)
    # Use case journey
    card(slide, 0.65, 3.77, 12.03, 1.28, PALE, RGBColor(193, 234, 231))
    add_text(slide, "核心用例", 0.9, 3.94, 1.0, 0.3, 13, TEAL, True)
    journey = [
        ("新建会话", MINT), ("多轮描述", BLUE), ("RAG 检索", PURPLE),
        ("安全门禁", RED), ("建议卡片", GREEN), ("历史沉淀", TEAL)
    ]
    for i, (text, color) in enumerate(journey):
        x = 1.95 + i * 1.72
        card(slide, x, 3.96, 1.32, 0.48, WHITE, color)
        add_text(slide, text, x, 3.96, 1.32, 0.48, 10, color, True, PP_ALIGN.CENTER)
        if i < len(journey) - 1:
            line(slide, x + 1.33, 4.2, x + 1.64, 4.2, color, 1.5, True)
    # Boundary
    card(slide, 0.65, 5.35, 5.75, 1.24, WHITE, RGBColor(190, 234, 218))
    add_text(slide, "✓  范围内", 0.92, 5.56, 1.3, 0.28, 13, GREEN, True)
    add_text(slide, "问诊 · RAG · 门禁 · 就诊卡 · 知识库 · 多媒体", 0.92, 5.93, 5.0, 0.3, 11.5, TEXT)
    card(slide, 6.72, 5.35, 5.96, 1.24, WHITE, RGBColor(244, 207, 211))
    add_text(slide, "×  范围外", 6.99, 5.56, 1.3, 0.28, 13, RED, True)
    add_text(slide, "真实诊疗 · 电子处方 · 医保结算 · 自由开药", 6.99, 5.93, 5.0, 0.3, 11.5, TEXT)
    add_footer(slide, 3)
    return slide


def slide_features(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 4, "FEATURES", "功能全景：围绕一次问诊形成完整闭环", "不只是聊天窗口，而是输入、知识、安全、档案与沉淀协同的软件系统")
    features = [
        ("01", "多轮智能问诊", "连续追问症状、时长、程度等关键信息", MINT),
        ("02", "RAG 知识增强", "混合检索科普资料，回答附知识依据", BLUE),
        ("03", "双重安全门禁", "红旗急诊出口 + 药品白名单校验", RED),
        ("04", "结构化建议卡片", "诊断参考、处理建议、用药与免责声明", ORANGE),
        ("05", "多就诊卡管理", "维护基础信息、既往史、过敏史并关联问诊", PURPLE),
        ("06", "历史会话管理", "打开、重命名、单删与多选批量删除", TEAL),
        ("07", "知识库前台", "分类浏览、关键词检索与教学资料补录", GREEN),
        ("08", "语音与图片输入", "语音转写、体征图片观察，仍统一经过门禁", BLUE),
    ]
    for i, (num, title, desc, color) in enumerate(features):
        col, row = i % 4, i // 4
        x = 0.62 + col * 3.15
        y = 2.07 + row * 2.15
        card(slide, x, y, 2.82, 1.77, WHITE, LINE)
        card(slide, x, y, 0.57, 1.77, color, color)
        add_text(slide, num, x + 0.03, y + 0.08, 0.5, 0.34, 10, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, title, x + 0.78, y + 0.23, 1.78, 0.38, 14, NAVY, True)
        add_text(slide, desc, x + 0.78, y + 0.78, 1.78, 0.67, 10, TEXT, valign=MSO_ANCHOR.TOP)
    add_text(
        slide, "贯穿能力：SSE 流式反馈 · 过期会话自动重建 · 检索组件优雅降级",
        0.8, 6.45, 11.7, 0.34, 11, TEAL, True, PP_ALIGN.CENTER
    )
    add_footer(slide, 4)
    return slide


def slide_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 5, "ARCHITECTURE", "总体架构：前后端分离，安全能力落在 Java", "必开组件精简；检索与图谱可缺席降级")
    # Layers
    layers = [
        (2.05, "表现层 · React + Vite", ["问诊页", "就诊卡", "知识库", "历史侧栏", "SSE 卡片"], TEAL),
        (3.05, "应用层 · Spring Boot", ["REST / SSE", "ConsultOrchestrator", "Profile", "Media", "Knowledge"], BLUE),
        (4.05, "领域层 · domain", ["SafetyGate", "白名单", "红旗否定", "会话记忆", "就诊卡规则"], PURPLE),
    ]
    for y, title, items, color in layers:
        card(slide, 0.72, y, 8.45, 0.78, WHITE, color)
        card(slide, 0.72, y, 2.05, 0.78, color, color)
        add_text(slide, title, 0.86, y, 1.77, 0.78, 11.5, WHITE, True, PP_ALIGN.CENTER)
        for i, item in enumerate(items):
            x = 2.98 + i * 1.17
            add_label(slide, item, x, y + 0.21, 1.03, PALE, TEAL)
    # Data and external, right side
    card(slide, 9.48, 2.05, 3.18, 2.78, WHITE, LINE)
    add_text(slide, "外部能力", 9.78, 2.28, 1.5, 0.3, 15, NAVY, True)
    ext = [
        ("LLM", "追问与生成", MINT),
        ("RAG", "向量 / 关键词 / 图谱", BLUE),
        ("多媒体", "语音转写 / 图片观察", ORANGE),
    ]
    for i, (name, desc, color) in enumerate(ext):
        y = 2.82 + i * 0.63
        circle(slide, 9.8, y, 0.34, color)
        add_text(slide, name, 10.3, y - 0.01, 0.7, 0.28, 11, NAVY, True)
        add_text(slide, desc, 11.0, y - 0.01, 1.35, 0.28, 9.5, TEXT)
    # Database bar
    card(slide, 0.72, 5.14, 11.94, 0.74, RGBColor(240, 249, 252), RGBColor(183, 220, 234))
    add_text(slide, "MySQL 8", 0.98, 5.14, 1.4, 0.74, 14, BLUE, True, PP_ALIGN.CENTER)
    add_text(slide, "会话 · 消息 · 就诊卡 · 病/症/药/红旗 · 知识底账", 2.35, 5.14, 6.2, 0.74, 12, TEXT)
    add_label(slide, "必开", 10.95, 5.33, 0.7, RGBColor(231, 247, 240), GREEN)
    # Ports
    add_text(slide, "开发部署", 0.78, 6.15, 1.0, 0.28, 11, TEAL, True)
    add_text(slide, "5173  前端", 1.85, 6.15, 1.2, 0.28, 10, TEXT)
    add_text(slide, "8082  后端", 3.2, 6.15, 1.2, 0.28, 10, TEXT)
    add_text(slide, "3306  MySQL", 4.55, 6.15, 1.4, 0.28, 10, TEXT)
    add_text(slide, "8000/8001  向量（可选）", 6.15, 6.15, 2.1, 0.28, 10, MUTED)
    add_text(slide, "7687  图谱（可选）", 8.55, 6.15, 1.8, 0.28, 10, MUTED)
    add_footer(slide, 5)
    return slide


def slide_pipeline(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 6, "CORE FLOW", "主链路：LLM → RAG → 生成 → 门禁", "串行而非三路并行；门禁是输出前不可绕过的最后一道质检")
    stages = [
        ("01", "用户输入", "症状 / 时长 / 发热\n+ 可选就诊卡", TEAL),
        ("02", "LLM 理解", "决定追问或作答\n保持同一会话上下文", BLUE),
        ("03", "RAG 工具", "向量 + 关键词 + 图谱\n融合相关资料", PURPLE),
        ("04", "LLM 草稿", "诊断参考 / 处理建议\n拟推荐药品", ORANGE),
        ("05", "Java 门禁", "红旗？白名单？\n表述是否安全？", RED),
        ("06", "结构化卡片", "通过则返回\n不通过改写就医建议", GREEN),
    ]
    start_x = 0.55
    for i, (num, title, desc, color) in enumerate(stages):
        x = start_x + i * 2.12
        card(slide, x, 2.55, 1.78, 2.16, WHITE, color)
        circle(slide, x + 0.17, 2.72, 0.42, color)
        add_text(slide, num, x + 0.17, 2.72, 0.42, 0.42, 9, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, title, x + 0.18, 3.25, 1.42, 0.38, 15, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.18, 3.78, 1.42, 0.66, 10, TEXT, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)
        if i < len(stages) - 1:
            line(slide, x + 1.8, 3.62, x + 2.04, 3.62, color, 2, True)
    # Safety branch
    line(slide, 9.94, 4.72, 9.94, 5.18, RED, 1.8, True)
    card(slide, 7.58, 5.2, 4.75, 0.78, RGBColor(255, 244, 245), RGBColor(244, 199, 205))
    add_text(slide, "红旗命中：不诊断、不开药，只导向急诊 / 120", 7.82, 5.2, 4.25, 0.78, 12, RED, True, PP_ALIGN.CENTER)
    # Degradation strip
    card(slide, 0.55, 5.44, 6.55, 0.92, PALE, RGBColor(191, 232, 230))
    add_text(slide, "优雅降级", 0.83, 5.62, 1.0, 0.28, 12, TEAL, True)
    add_text(slide, "向量/图谱可缺席 → 关键词/MySQL 兜底；系统变笨，但不能变危险", 1.92, 5.55, 4.85, 0.45, 10.5, TEXT)
    add_footer(slide, 6)
    return slide


def add_screenshot_card(slide, path, x, y, w, h, label, accent=MINT):
    card(slide, x, y, w, h, WHITE, LINE)
    add_picture_crop(slide, path, x + 0.08, y + 0.08, w - 0.16, h - 0.62)
    add_label(slide, label, x + 0.16, y + h - 0.45, w - 0.32, PALE, accent)


def slide_demo(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 7, "DEMO", "功能演示：从问诊到档案与知识沉淀", "建议现场依次演示普通问诊 → 红旗 → 否定纠错 → 就诊卡 → 知识库")
    add_screenshot_card(slide, ASK_SHOT, 0.55, 2.08, 6.15, 4.62, "① 问诊主界面 + 历史侧栏", TEAL)
    add_screenshot_card(slide, RED_FLAG_SHOT, 6.92, 2.08, 5.83, 1.52, "② 红旗只导向急诊", RED)
    add_screenshot_card(slide, PROFILE_SHOT, 6.92, 3.83, 2.79, 2.87, "③ 多就诊卡", BLUE)
    add_screenshot_card(slide, KB_SHOT, 9.96, 3.83, 2.79, 2.87, "④ 知识库", GREEN)
    add_footer(slide, 7)
    return slide


def slide_bug_evidence(prs, index, title, subtitle, image, condition, expected, actual, severity, accent):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, index, "BUG EVIDENCE", title, subtitle)
    card(slide, 0.58, 2.05, 7.55, 4.63, WHITE, LINE)
    add_picture_contain(slide, image, 0.68, 2.15, 7.35, 4.06)
    add_label(slide, f"同一条件重复 3 次 · 复现率 3/3 · {severity}", 0.86, 6.28, 6.95, PALE, accent)

    detail_x = 8.42
    blocks = [
        ("激活条件", condition, TEAL),
        ("预期结果", expected, GREEN),
        ("实际结果", actual, accent),
    ]
    for i, (label, body, color) in enumerate(blocks):
        y = 2.05 + i * 1.48
        card(slide, detail_x, y, 4.28, 1.25, WHITE, color)
        add_label(slide, label, detail_x + 0.18, y + 0.16, 1.02, color, WHITE)
        add_text(slide, body, detail_x + 1.4, y + 0.13, 2.58, 0.94, 10.5, TEXT, valign=MSO_ANCHOR.TOP)
    add_footer(slide, index, "青禾门诊 · Bug 激活条件与复现证据")
    return slide


def slide_bug_1(prs):
    return slide_bug_evidence(
        prs,
        8,
        "Bug 1｜正在读文章时，搜索框不生效",
        "正常打开一篇科普后，左侧搜索框仍然能输入，但结果不会替换当前文章",
        BUG_1_SHOT,
        "打开《腹泻先补液》，在左侧搜索框输入“流感”。",
        "退出当前文章，展示流感相关知识条目。",
        "搜索框已显示“流感”，页面却仍停留在腹泻文章。",
        "中等 · 页面状态",
        ORANGE,
    )


def slide_bug_2(prs):
    return slide_bug_evidence(
        prs,
        9,
        "Bug 2｜感冒科普被分进用药指南",
        "分类程序先匹配“抗生素”等词，文章还没按疾病主题归类就被改了类别",
        BUG_2_SHOT,
        "进入知识库，直接点击左侧“用药指南”。",
        "《普通感冒和流感不要混为一谈》应出现在“常见疾病”。",
        "它出现在“用药指南”中；“常见疾病”里反而找不到。",
        "中等 · 分类逻辑",
        ORANGE,
    )


def slide_bug_3(prs):
    return slide_bug_evidence(
        prs,
        10,
        "Bug 3｜新会话首条消息未应用已关联就诊卡",
        "界面显示“已关联”，但创建会话后没有把所选 profileId 同步到后端",
        BUG_3_SHOT,
        "开启新对话 → 选择“李四女的就诊卡” → 首条发送“我叫什么名字？”。",
        "新会话写入所选卡片，并根据就诊卡回答姓名。",
        "页面仍显示已关联，后端 profileLinked=false，回复追问哪里不舒服。",
        "严重 · 状态同步",
        RED,
    )


def slide_safety_test(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 8, "QUALITY", "安全与测试：模型负责表达，规则负责底线", "自动化 + 场景演示 + 缺陷回归三层覆盖")
    # Three safeguards
    guards = [
        ("红旗门禁", "胸痛、大出血、意识异常等\n仅给急诊建议", RED),
        ("药品白名单", "模型草稿中的药品必须命中\n库内条目才能展示", GREEN),
        ("否定语处理", "“胸口疼不存在了”可解除\n避免反复错误拦截", BLUE),
    ]
    for i, (title, desc, color) in enumerate(guards):
        x = 0.65 + i * 4.12
        card(slide, x, 2.08, 3.72, 1.48, WHITE, color)
        shape(slide, MSO_SHAPE.PENTAGON, x + 0.22, 2.35, 0.64, 0.7, color, color)
        add_text(slide, title, x + 1.02, 2.23, 2.3, 0.32, 15, NAVY, True)
        add_text(slide, desc, x + 1.02, 2.66, 2.32, 0.58, 10.5, TEXT, valign=MSO_ANCHOR.TOP)
    # Metrics
    metrics = [
        ("11", "类自动化测试", "领域规则 + Spring 流程"),
        ("9", "个场景用例", "正常 / 安全 / 异常"),
        ("3", "个关键 Bug", "均纳入回归设计"),
    ]
    for i, (value, title, sub) in enumerate(metrics):
        x = 0.65 + i * 2.75
        card(slide, x, 3.92, 2.5, 1.42, PALE, RGBColor(191, 232, 230))
        add_text(slide, value, x + 0.18, 4.05, 0.72, 0.55, 28, TEAL, True, PP_ALIGN.CENTER)
        add_text(slide, title, x + 0.94, 4.06, 1.32, 0.3, 12, NAVY, True)
        add_text(slide, sub, x + 0.94, 4.48, 1.3, 0.32, 9.5, MUTED)
    # Bugs
    card(slide, 8.98, 3.92, 3.7, 2.2, WHITE, LINE)
    add_text(slide, "验收中真实修复", 9.25, 4.15, 2.8, 0.3, 14, NAVY, True)
    bug_items = [
        "过期会话号 → 自动重建",
        "红旗否定失效 → 后置否定窗口",
        "多卡唯一索引 → 启动时修复 Schema",
    ]
    add_bullet_list(slide, bug_items, 9.25, 4.63, 2.9, 1.2, 10.5, gap=0.42)
    add_text(slide, "测试原则：变笨可以，变危险不行", 0.78, 5.75, 7.5, 0.38, 15, RED, True)
    add_footer(slide, 8)
    return slide


def slide_stack_roles(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 11, "DELIVERY", "技术选型与小组分工", "选型围绕可维护、可演示、可降级；成员姓名可直接在此页替换")
    stacks = [
        ("前端", "React · Vite · TypeScript\nSSE 消费 · 响应式 UI", TEAL),
        ("后端", "Java 21 · Spring Boot 3.4\nREST/SSE · 编排 · 门禁", BLUE),
        ("数据与 AI", "MySQL 8 · Chroma · Neo4j\nDeepSeek / OpenAI 兼容", PURPLE),
        ("测试与文档", "JUnit · Maven Surefire\n场景用例 · Bug 回归", ORANGE),
    ]
    for i, (title, desc, color) in enumerate(stacks):
        x = 0.65 + i * 3.05
        card(slide, x, 2.05, 2.75, 1.45, WHITE, color)
        card(slide, x, 2.05, 2.75, 0.42, color, color)
        add_text(slide, title, x, 2.05, 2.75, 0.42, 12, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.2, 2.68, 2.35, 0.58, 10.5, TEXT, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)
    # Division table
    card(slide, 0.65, 3.85, 12.03, 2.34, WHITE, LINE)
    add_text(slide, "角色", 0.92, 4.04, 1.1, 0.3, 11, MUTED, True)
    add_text(slide, "负责内容", 2.15, 4.04, 6.4, 0.3, 11, MUTED, True)
    add_text(slide, "成员", 9.35, 4.04, 2.3, 0.3, 11, MUTED, True)
    line(slide, 0.88, 4.42, 12.35, 4.42, LINE, 1)
    rows = [
        ("产品 / 文档", "范围、报告、PPT、演示剧本"),
        ("后端 / 检索", "问诊编排、RAG、门禁、会话、就诊卡 API"),
        ("前端 / 测试", "三页 UI、历史交互、多媒体、用例与回归"),
    ]
    for i, (role, work) in enumerate(rows):
        y = 4.56 + i * 0.48
        add_text(slide, role, 0.92, y, 1.18, 0.32, 10.5, NAVY, True)
        add_text(slide, work, 2.15, y, 6.7, 0.32, 10.5, TEXT)
        add_text(slide, "________________", 9.35, y, 2.0, 0.32, 10.5, MUTED)
    add_footer(slide, 11)
    return slide


def slide_summary(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, 12, "WRAP-UP", "从“会聊天”到“可运行、可解释、可回归”", "项目完成了完整软件工程闭环，也通过边界测试发现了可稳定复现的问题")
    # Main completion
    card(slide, 0.65, 2.05, 5.1, 4.3, TEAL, TEAL)
    add_text(slide, "已完成", 0.98, 2.35, 1.25, 0.34, 14, PALE, True)
    add_text(slide, "一个可演示的\n智能问诊软件系统", 0.98, 2.9, 3.9, 1.1, 27, WHITE, True, valign=MSO_ANCHOR.TOP)
    done = ["问诊 + RAG + 门禁", "就诊卡 + 历史 + 知识库", "多媒体输入 + 优雅降级", "测试用例 + Bug 回归 + 项目报告"]
    add_bullet_list(slide, done, 1.02, 4.25, 3.95, 1.8, 12, WHITE, RGBColor(179, 239, 234), 0.43)
    # Limits and future
    card(slide, 6.08, 2.05, 6.6, 1.85, WHITE, LINE)
    add_text(slide, "当前不足", 6.4, 2.3, 1.4, 0.32, 15, NAVY, True)
    add_bullet_list(
        slide,
        ["知识覆盖仍以教学资料包为主", "前端 E2E 与高并发压测尚未补齐", "多用户权限未纳入课程演示范围"],
        6.42, 2.74, 5.65, 0.98, 10.5, RED, gap=0.33
    )
    card(slide, 6.08, 4.18, 6.6, 2.17, WHITE, LINE)
    add_text(slide, "下一步", 6.4, 4.43, 1.4, 0.32, 15, NAVY, True)
    futures = [
        ("扩知识", "更多病种与权威资料"),
        ("补测试", "Playwright + 性能冒烟"),
        ("强解释", "检索证据与降级状态可视化"),
    ]
    for i, (title, desc) in enumerate(futures):
        y = 4.92 + i * 0.43
        add_label(slide, title, 6.42, y, 0.85, PALE, TEAL)
        add_text(slide, desc, 7.48, y, 4.3, 0.36, 10.5, TEXT)
    add_footer(slide, 12)
    return slide


def slide_thanks(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, TEAL)
    add_brand(slide, inverse=True)
    add_text(slide, "谢谢聆听", 0.72, 1.36, 5.2, 0.85, 44, WHITE, True, valign=MSO_ANCHOR.TOP)
    add_text(slide, "青禾门诊 · 智能问诊系统", 0.77, 2.34, 4.8, 0.42, 19, PALE)
    line(slide, 0.78, 3.05, 5.55, 3.05, RGBColor(92, 194, 197), 1.2)
    add_text(slide, "AI 让医疗更有温度", 0.78, 3.32, 4.5, 0.4, 16, WHITE, True)
    add_text(slide, "教学演示 · 不替代面诊 · 不生成真实处方", 0.78, 3.84, 5.2, 0.34, 11, RGBColor(198, 240, 238))
    circle(slide, 7.45, 0.8, 4.9, RGBColor(243, 253, 252), RGBColor(177, 230, 227))
    slide.shapes.add_picture(str(ROBOT), Inches(7.85), Inches(1.18), Inches(4.1), Inches(4.1))
    card(slide, 7.75, 5.48, 4.25, 0.7, RGBColor(14, 137, 151), RGBColor(55, 178, 184))
    add_text(slide, "欢迎现场交流与提问", 7.75, 5.48, 4.25, 0.7, 15, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, "13", 12.22, 7.08, 0.4, 0.22, 9, PALE, True, PP_ALIGN.RIGHT)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "青禾门诊智能问诊系统·课堂汇报"
    prs.core_properties.subject = "软件工程课程小组项目"
    prs.core_properties.author = "青禾门诊项目组"
    prs.core_properties.keywords = "智能问诊,RAG,安全门禁,软件工程"

    slide_cover(prs)
    slide_background(prs)
    slide_requirements(prs)
    slide_features(prs)
    slide_architecture(prs)
    slide_pipeline(prs)
    slide_demo(prs)
    slide_bug_1(prs)
    slide_bug_2(prs)
    slide_bug_3(prs)
    slide_stack_roles(prs)
    slide_summary(prs)
    slide_thanks(prs)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
